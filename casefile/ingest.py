"""Find files, pull plain text out of them, and cut the text into model-sized chunks.

Every chunk remembers where it came from so any extracted fact can be cited back to
a file and a page (PDF) or line (everything else).

Containers (an email with attachments, an .mbox mailbox, an Outlook .msg) become
several Documents: the message itself plus one per attachment, named
"<container> > <attachment>", so each attachment is read and cited like a file.
"""

from __future__ import annotations

import bisect
import codecs
import contextlib
import datetime as dt
import email
import email.message
import email.policy
import email.utils
import hashlib
import io
import json
import mailbox
import mimetypes
import os
import re
import shutil
import subprocess
import tempfile
import threading
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Callable, Iterator
from xml.etree import ElementTree

TEXT_EXTS = {".txt", ".md", ".markdown", ".log", ".csv", ".tsv", ".json", ".xml", ".yaml", ".yml"}
HTML_EXTS = {".html", ".htm"}
SUPPORTED_EXTS = TEXT_EXTS | HTML_EXTS | {
    ".pdf", ".docx", ".doc", ".rtf", ".odt", ".xlsx", ".pptx", ".eml", ".mbox", ".msg",
}

SKIPPED = "skipped: unsupported file type"
NO_PDF_TEXT_NO_OCR = (
    "no text found (scanned image? needs OCR: install Tesseract and casefile will read it)"
)
NO_PDF_TEXT = "no text found (scanned image, and OCR found no words)"
PDF_PASSWORD = "password protected: remove the password to read it"
OLD_DOC = "old Word .doc: open it and save as .docx to read it"
# An attachment inside an attachment inside ... stops somewhere; real mail rarely nests past 2.
MAX_NESTING = 5
# Bump when PDF reading changes (text extraction, OCR, repeated-line removal): cached
# extracted text is keyed on it, so old entries are ignored rather than reused.
EXTRACT_VERSION = "1"
# A line this long that recurs this often in one PDF is a page header, footer or a quoted
# copy of an earlier email: it is kept the first time and dropped after that.
REPEAT_MIN_CHARS = 20
REPEAT_MIN_COUNT = 3

Source = Path | bytes  # a file on disk, or the bytes of an attachment
Progress = Callable[[str], None]  # what is being read right now, for a progress display
Remember = Callable[[str], None]  # a cache key this run used, so the caller can delete it later


@dataclass
class Document:
    path: str  # relative to the scanned folder; attachments are "<container> > <name>"
    text: str = ""
    page_starts: list[int] | None = None  # PDF only: text offset where each page begins
    status: str = "ok"
    repeats_removed: int = 0  # PDF only: lines dropped as repeated headers, footers or quoted copies
    _line_starts: list[int] = field(default_factory=list, repr=False)

    def location(self, offset: int) -> str:
        """Human citation for a character offset: 'p. 4' for PDFs, 'line 120' otherwise."""
        if self.page_starts is not None:
            return f"p. {bisect.bisect_right(self.page_starts, offset)}"
        if not self._line_starts:
            self._line_starts = [0] + [i + 1 for i, ch in enumerate(self.text) if ch == "\n"]
        return f"line {bisect.bisect_right(self._line_starts, offset)}"


@dataclass
class Chunk:
    doc_index: int
    index: int  # position within its document
    start: int  # offset into Document.text
    text: str


def scan_folder(
    root: Path, cache_dir: Path | None = None, remember: Remember | None = None, progress: Progress | None = None
) -> list[Document]:
    """Read every file under root (skipping hidden files/dirs), in stable path order.

    With cache_dir, the text read out of each PDF (OCR included) is kept there, keyed on the
    file's bytes, so a rerun doesn't read or OCR it again; every key used is passed to remember."""
    docs = []
    for path in sorted(p for p in root.rglob("*") if p.is_file()):
        rel = path.relative_to(root)
        if any(part.startswith(".") for part in rel.parts):
            continue
        if progress:
            progress(f"Reading {rel.as_posix()}")
        docs.extend(read_document(path, rel.as_posix(), cache_dir, remember, progress))
    return docs


def read_document(
    path: Path, rel: str, cache_dir: Path | None = None, remember: Remember | None = None, progress: Progress | None = None
) -> list[Document]:
    """Read one file. Most files give one Document; mail containers give one per
    message and attachment, the container's own text first."""
    return _read(path, Path(rel).suffix.lower(), rel, 0, cache_dir, remember, progress)


def _read(
    src: Source, ext: str, rel: str, depth: int,
    cache_dir: Path | None = None, remember: Remember | None = None, progress: Progress | None = None,
) -> list[Document]:
    if ext not in SUPPORTED_EXTS:
        return [Document(rel, status=SKIPPED)]
    try:
        if ext == ".eml":
            msg = email.message_from_bytes(_bytes(src), policy=email.policy.default)
            return _email_docs(msg, rel, depth)
        if ext == ".mbox":
            return _read_mbox(src, rel, depth)
        if ext == ".msg":
            return _msg_docs(_ole_message(_bytes(src)), rel, depth)
        if ext == ".pdf":
            doc = _read_pdf_cached(_bytes(src), rel, cache_dir, remember, progress)
        elif ext == ".doc":
            doc = _read_doc(src, rel)
        else:
            doc = Document(rel, _plain_text(src, ext))
    except Exception as exc:  # one unreadable file must not stop a scan of thousands
        return [Document(rel, status=f"error: {type(exc).__name__}: {exc}")]
    if doc.status == "ok" and not doc.text.strip():
        doc.status = "empty"
    return [doc]


def _plain_text(src: Source, ext: str) -> str:
    if ext == ".docx":
        return _read_docx(_bytes(src))
    if ext == ".rtf":
        return _rtf_to_text(_bytes(src))
    if ext == ".odt":
        return _read_odt(_bytes(src))
    if ext == ".xlsx":
        return _read_xlsx(_bytes(src))
    if ext == ".pptx":
        return _read_pptx(_bytes(src))
    if ext in HTML_EXTS:
        return _html_to_text(_decode(_bytes(src)))
    return _decode(_bytes(src))


def _bytes(src: Source) -> bytes:
    return src.read_bytes() if isinstance(src, Path) else src


@contextlib.contextmanager
def _on_disk(src: Source, ext: str) -> Iterator[Path]:
    """A path for tools that only take files: the file itself, or a temp copy of an attachment."""
    if isinstance(src, Path):
        yield src
        return
    with tempfile.TemporaryDirectory(prefix="casefile-") as tmp:
        path = Path(tmp) / f"attachment{ext}"
        path.write_bytes(src)
        yield path


def chunk_document(doc: Document, doc_index: int, size: int, overlap: int) -> list[Chunk]:
    """Split on line boundaries into chunks that overlap, so an event straddling a
    boundary is seen whole by at least one chunk.

    size and overlap are measured in Latin characters (about 4 per model token). Chinese,
    Japanese, Korean and Thai text runs at about 1 character per token, so those characters
    count 4 each: every chunk then costs the model about size/4 tokens whatever the script."""
    text = doc.text
    chunks: list[Chunk] = []
    start = 0
    while start < len(text):
        end = _reach(text, start, size)
        if end < len(text):
            cut = text.rfind("\n", start + (end - start) // 2, end)
            if cut != -1:
                end = cut + 1
        piece = text[start:end]
        if piece.strip():
            chunks.append(Chunk(doc_index, len(chunks), start, piece))
        if end >= len(text):
            break
        next_start = max(_back(text, end, overlap), start + 1)
        nl = text.find("\n", next_start, end)
        # Start the next chunk on a line boundary, unless that would leave no overlap at all.
        start = nl + 1 if nl != -1 and nl + 1 < end else next_start
    return chunks


# Scripts that tokenizers split into about one token per character (vs ~4 Latin characters):
# CJK ideographs and punctuation, kana, Hangul, full-width forms, and Thai/Lao/Myanmar/Khmer.
_WIDE = re.compile(
    "[\u0e00-\u0eff\u1000-\u109f\u1100-\u11ff\u1780-\u17ff\u2e80-\u303f\u3040-\u30ff"
    "\u3130-\u318f\u31f0-\u31ff\u3400-\u4dbf\u4e00-\u9fff\ua960-\ua97f\uac00-\ud7ff"
    "\uf900-\ufaff\uff00-\uffef\U00020000-\U0003ffff]"
)
_WIDE_EXTRA = 3  # a wide character costs 1 + 3 = 4 Latin characters


def _cost(text: str, start: int, end: int) -> int:
    return (end - start) + _WIDE_EXTRA * len(_WIDE.findall(text, start, end))


def _reach(text: str, start: int, budget: int) -> int:
    """Furthest end (> start) such that text[start:end] costs about budget or less."""
    end = min(start + budget, len(text))
    while end > start + 1:
        cost = _cost(text, start, end)
        if cost <= budget:
            break
        # Shrink in proportion to the overshoot; always make progress.
        end = min(end - 1, start + max(1, (end - start) * budget // cost))
    return max(end, start + 1)


def _back(text: str, end: int, budget: int) -> int:
    """Earliest start (<= end) such that text[start:end] costs about budget or less."""
    start = max(end - budget, 0)
    while start < end:
        cost = _cost(text, start, end)
        if cost <= budget:
            break
        start = max(start + 1, end - max(0, (end - start) * budget // cost))
    return start


def _decode(data: bytes) -> str:
    # cp1252 accepts almost any bytes, so UTF-16/32 (Windows "Unicode" text) must be
    # recognised by its byte-order mark first or it turns into NUL-riddled garbage.
    if data.startswith((codecs.BOM_UTF32_LE, codecs.BOM_UTF32_BE)):
        return data.decode("utf-32", errors="replace")
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16", errors="replace")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


# ---------------------------------------------------------------------------- PDF


def _read_pdf_cached(
    data: bytes, rel: str, cache_dir: Path | None, remember: Remember | None, progress: Progress | None
) -> Document:
    """_read_pdf, remembered on disk: reading a scanned PDF costs about a second a page, so
    the result is kept under a hash of the file's bytes (and of how it was read)."""
    if cache_dir is None:
        return _read_pdf(data, rel, progress)
    key = "doc-" + hashlib.sha256(
        b"\0".join((EXTRACT_VERSION.encode(), b"ocr" if _tesseract() else b"", data))
    ).hexdigest()
    cached = cache_dir / f"{key}.json"
    try:
        saved = json.loads(cached.read_text())
        doc = Document(rel, saved["text"], saved["page_starts"], saved["status"], saved["repeats_removed"])
    except (OSError, ValueError, KeyError, TypeError):
        doc = _read_pdf(data, rel, progress)
        cache_dir.mkdir(parents=True, exist_ok=True)
        tmp = cache_dir / f"{key}.{threading.get_ident()}.tmp"
        tmp.write_text(json.dumps(
            {"text": doc.text, "page_starts": doc.page_starts, "status": doc.status, "repeats_removed": doc.repeats_removed}
        ))
        tmp.replace(cached)  # atomic: an interrupted run never leaves a half-written file
    if remember:
        remember(key)
    return doc


def _read_pdf(data: bytes, rel: str, progress: Progress | None = None) -> Document:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        # Many bank/court PDFs are "secured" with an empty open password: they only
        # restrict editing or printing, and anyone can read them.
        if not reader.decrypt(""):
            return Document(rel, status=PDF_PASSWORD)
    texts = [(page.extract_text() or "").strip() for page in reader.pages]
    blank = [i for i, t in enumerate(texts) if not t]
    tesseract = _tesseract()
    if blank and tesseract:
        # Scanned pages carry only an image; OCR just those so typed pages keep exact text.
        def on_page(done: int) -> None:
            if progress:
                progress(f"Reading {rel}: scanned page {done} of {len(blank)}")

        try:
            ocr = _ocr_pages(data, blank, tesseract, on_page)
        except Exception:  # OCR is a bonus: a failure must not lose the typed pages
            ocr = {}
        for i, words in ocr.items():
            texts[i] = words.strip()
    texts, repeats = _drop_repeated_lines(texts)
    parts, starts, offset = [], [], 0
    for page_text in texts:
        starts.append(offset)
        parts.append(page_text + "\n\n")
        offset += len(page_text) + 2
    doc = Document(rel, "".join(parts), page_starts=starts, repeats_removed=repeats)
    if not doc.text.strip():
        doc.status = NO_PDF_TEXT if tesseract else NO_PDF_TEXT_NO_OCR
    return doc


_QUOTE_PREFIX = re.compile(r"^(?:\s*>)+\s?")


def _drop_repeated_lines(pages: list[str]) -> tuple[list[str], int]:
    """Keep the first copy of any line that recurs REPEAT_MIN_COUNT times or more in the
    document and drop the rest: page headers and footers, and the quoted copies of earlier
    messages that every reply in a printed email chain carries. The model then reads each
    passage once, and the first copy keeps its date and context. Only long lines count,
    compared without case, surrounding space or '>' quote marks, so short lines that
    legitimately recur (dates, amounts, names) are left alone."""
    counts: Counter[str] = Counter()
    for page in pages:
        for line in page.split("\n"):
            n = _repeat_key(line)
            if n:
                counts[n] += 1
    if not any(c >= REPEAT_MIN_COUNT for c in counts.values()):
        return pages, 0
    seen: set[str] = set()
    out, dropped = [], 0
    for page in pages:
        kept = []
        for line in page.split("\n"):
            n = _repeat_key(line)
            if n and counts[n] >= REPEAT_MIN_COUNT:
                if n in seen:
                    dropped += 1
                    continue
                seen.add(n)
            kept.append(line)
        out.append("\n".join(kept).strip())
    return out, dropped


def _repeat_key(line: str) -> str:
    n = " ".join(_QUOTE_PREFIX.sub("", line).split()).casefold()
    return n if len(n) >= REPEAT_MIN_CHARS else ""


_BIN_DIRS = ("/opt/homebrew/bin", "/usr/local/bin", "/usr/bin")


def _find_tool(name: str, extra: tuple[str, ...] = ()) -> str | None:
    """A command-line program, if installed. Apps started from the macOS Finder don't get
    the shell's PATH, so the usual install folders are checked as well."""
    return shutil.which(name) or shutil.which(name, path=os.pathsep.join(_BIN_DIRS + extra))


def _tesseract() -> str | None:
    return _find_tool("tesseract", (r"C:\Program Files\Tesseract-OCR",))


_OCR_DPI = 300
_OCR_MAX_SIDE = 7000  # pixels; keeps a poster-sized page from needing gigabytes
_NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0)  # Windows: no console flash


def _ocr_pages(data: bytes, pages: list[int], tesseract: str, on_page: Callable[[int], None] | None = None) -> dict[int, str]:
    """Render each listed page to a grayscale image and read it with the tesseract CLI.
    on_page(n) is called after each page, n pages done."""
    import pypdfium2 as pdfium

    out: dict[int, str] = {}
    pdf = pdfium.PdfDocument(data)
    try:
        with tempfile.TemporaryDirectory(prefix="casefile-ocr-") as tmp:
            for done, i in enumerate(pages, 1):
                page = pdf[i]
                width, height = page.get_size()
                scale = min(_OCR_DPI / 72, _OCR_MAX_SIDE / max(width, height, 1))
                bitmap = page.render(scale=scale, grayscale=True)
                image = Path(tmp) / f"page{i}.pgm"
                image.write_bytes(_pgm(bitmap))
                try:
                    result = subprocess.run(
                        [tesseract, str(image), "stdout"],
                        capture_output=True,
                        timeout=300,
                        creationflags=_NO_WINDOW,
                    )
                except subprocess.TimeoutExpired:
                    result = None  # one stuck page leaves that page blank, not the whole PDF
                if result is not None and result.returncode == 0:
                    out[i] = result.stdout.decode("utf-8", errors="replace")
                image.unlink()
                bitmap.close()
                page.close()
                if on_page:
                    on_page(done)
    finally:
        pdf.close()
    return out


def _pgm(bitmap) -> bytes:
    """A PDFium bitmap as a binary PGM image (tesseract reads it; no imaging library needed)."""
    w, h, stride, channels = bitmap.width, bitmap.height, bitmap.stride, bitmap.n_channels
    raw = bytes(bitmap.buffer)
    # Gray bitmaps have one byte per pixel; for colour ones the green byte is gray enough.
    first = 0 if channels == 1 else 1
    rows = b"".join(raw[y * stride + first : y * stride + w * channels : channels] for y in range(h))
    return b"P5\n%d %d\n255\n" % (w, h) + rows


# ------------------------------------------------------------------ Office formats

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _read_docx(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        root = ElementTree.fromstring(zf.read("word/document.xml"))
    paragraphs = []
    for para in root.iter(f"{_W}p"):
        bits = []
        for node in para.iter():
            if node.tag == f"{_W}t" and node.text:
                bits.append(node.text)
            elif node.tag == f"{_W}tab":
                bits.append("\t")
            elif node.tag in (f"{_W}br", f"{_W}cr"):
                bits.append("\n")
        paragraphs.append("".join(bits))
    return "\n".join(paragraphs)


def _read_doc(src: Source, rel: str) -> Document:
    """Old binary Word files: macOS's textutil or antiword when present."""
    textutil, antiword = _find_tool("textutil"), _find_tool("antiword")
    if not (textutil or antiword):
        return Document(rel, status=OLD_DOC)
    with _on_disk(src, ".doc") as path:
        if textutil:
            cmd = [textutil, "-convert", "txt", "-encoding", "UTF-8", "-stdout", str(path)]
        else:
            cmd = [antiword, "-m", "UTF-8.txt", str(path)]
        result = subprocess.run(cmd, capture_output=True, timeout=120, creationflags=_NO_WINDOW)
    if result.returncode != 0:
        reason = result.stderr.decode("utf-8", errors="replace").strip().splitlines()
        detail = f" ({reason[-1]})" if reason else ""
        return Document(rel, status=f"error: could not read this .doc{detail}; {OLD_DOC}")
    return Document(rel, result.stdout.decode("utf-8", errors="replace").replace("\r\n", "\n"))


_TEXT_NS = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"


def _read_odt(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        root = ElementTree.fromstring(zf.read("content.xml"))
    blocks = (f"{_TEXT_NS}p", f"{_TEXT_NS}h")
    lines = []
    for el in root.iter():
        if el.tag in blocks:
            bits: list[str] = []
            _odt_inline(el, bits, blocks)
            lines.append("".join(bits))
    return "\n".join(lines)


def _odt_inline(el, bits: list[str], blocks: tuple[str, str]) -> None:
    if el.text:
        bits.append(el.text)
    for child in el:
        if child.tag == f"{_TEXT_NS}s":
            bits.append(" " * int(child.get(f"{_TEXT_NS}c", "1")))
        elif child.tag == f"{_TEXT_NS}tab":
            bits.append("\t")
        elif child.tag == f"{_TEXT_NS}line-break":
            bits.append("\n")
        elif child.tag not in blocks:  # nested paragraphs (notes, frames) get their own line
            _odt_inline(child, bits, blocks)
        if child.tail:
            bits.append(child.tail)


def _read_xlsx(data: bytes) -> str:
    """One line per non-empty row, cells tab-separated, under a line naming each sheet.
    Values only: formulas give their last saved result."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines = []
    try:
        for ws in wb.worksheets:
            lines.append(f"Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = [_cell_text(v) for v in row]
                while cells and not cells[-1]:
                    cells.pop()
                if cells:
                    lines.append("\t".join(cells))
            lines.append("")
    finally:
        wb.close()
    return "\n".join(lines)


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat(sep=" ")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())


_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def _read_pptx(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = _pptx_slide_order(zf)
        out = []
        for n, name in enumerate(names, 1):
            root = ElementTree.fromstring(zf.read(name))
            out.append(f"Slide {n}")
            for para in root.iter(f"{_A}p"):
                line = "".join(t.text or "" for t in para.iter(f"{_A}t"))
                if line.strip():
                    out.append(line)
            out.append("")
    return "\n".join(out)


def _pptx_slide_order(zf: zipfile.ZipFile) -> list[str]:
    """Slides in presentation order (the list in presentation.xml), falling back to file numbers."""
    files = set(zf.namelist())
    try:
        pres = ElementTree.fromstring(zf.read("ppt/presentation.xml"))
        rels = ElementTree.fromstring(zf.read("ppt/_rels/presentation.xml.rels"))
        targets = {r.get("Id"): r.get("Target", "") for r in rels.iter(f"{_REL}Relationship")}
        order = []
        for sld in pres.iter(f"{_P}sldId"):
            target = targets.get(sld.get(f"{_R}id"), "")
            name = target.lstrip("/") if target.startswith("/") else f"ppt/{target}"
            if name in files:
                order.append(name)
        if order:
            return order
    except (KeyError, ElementTree.ParseError):
        pass
    slides = [n for n in files if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
    return sorted(slides, key=lambda n: int(re.search(r"(\d+)\.xml$", n).group(1)))


# ---------------------------------------------------------------------------- RTF

# Groups whose text isn't document content (font and colour tables, metadata, pictures...).
_RTF_SKIP = {
    "fonttbl", "colortbl", "stylesheet", "info", "pict", "listtable", "listoverridetable",
    "revtbl", "rsidtbl", "generator", "xmlnstbl", "themedata", "colorschememapping",
    "datastore", "latentstyles", "filetbl", "fldinst", "objdata", "object", "nonshppict",
    "shpinst", "pgdsctbl", "mmathPr", "xmlopen", "falt", "panose", "bkmkstart", "bkmkend",
}
_RTF_CHARS = {
    "par": "\n", "line": "\n", "sect": "\n\n", "page": "\n\n", "row": "\n", "cell": "\t",
    "tab": "\t", "emdash": "\u2014", "endash": "\u2013", "lquote": "\u2018", "rquote": "\u2019",
    "ldblquote": "\u201c", "rdblquote": "\u201d", "bullet": "\u2022", "emspace": " ",
    "enspace": " ", "qmspace": " ",
}
_RTF_TOKEN = re.compile(
    r"\\([a-zA-Z]+)(-?\d+)? ?"  # control word with optional number; one space is its delimiter
    r"|\\'([0-9a-fA-F]{2})"  # a byte in the document's code page
    r"|\\([^a-zA-Z])"  # control symbol: \\ \{ \} \~ \- \_ \*
    r"|([{}])"
    r"|[\r\n]+"  # raw line breaks are not text in RTF
    r"|([^\\{}\r\n]+)"
)


def _rtf_to_text(data: bytes) -> str:
    if not data.lstrip().startswith(b"{\\rtf"):
        return _decode(data)  # a plain-text file saved with an .rtf name
    raw = data.decode("latin-1")
    codepage = "cp1252"
    m = re.search(r"\\ansicpg(\d+)", raw[:2000])
    if m:
        with contextlib.suppress(LookupError):
            codecs.lookup(f"cp{m.group(1)}")
            codepage = f"cp{m.group(1)}"
    out: list[str] = []
    pending = bytearray()  # \'hh bytes, decoded together so multi-byte code pages work
    stack: list[tuple[bool, int]] = []
    skip, uc = False, 1
    to_drop = 0  # fallback characters still to drop after a \uN
    group_start = False
    pos = 0

    def flush() -> None:
        if pending:
            out.append(pending.decode(codepage, errors="replace"))
            pending.clear()

    while pos < len(raw):
        m = _RTF_TOKEN.match(raw, pos)
        if m is None:  # a lone trailing backslash
            break
        pos = m.end()
        word, num, hexbyte, symbol, brace, text = m.groups()
        first, group_start = group_start, False
        if hexbyte is not None:
            if to_drop:
                to_drop -= 1
            elif not skip:
                pending.append(int(hexbyte, 16))
            continue
        flush()
        if brace == "{":
            stack.append((skip, uc))
            group_start = True
        elif brace == "}":
            if stack:
                skip, uc = stack.pop()
            to_drop = 0
        elif symbol is not None:
            if symbol == "*" and first:
                skip = True
            elif skip:
                pass
            elif symbol in "\\{}":
                out.append(symbol)
            elif symbol == "~":
                out.append("\u00a0")
            elif symbol == "_":
                out.append("-")
            # \- (optional hyphen) and others add nothing
        elif word is not None:
            if word in _RTF_SKIP:
                skip = True
            elif word == "bin" and num:  # raw binary data follows; never text
                pos += int(num)
            elif word == "uc" and num:
                uc = int(num)
            elif word == "u" and num:
                if not skip:
                    out.append(chr(int(num) % 65536))
                to_drop = uc
            elif not skip and word in _RTF_CHARS:
                out.append(_RTF_CHARS[word])
        elif text is not None:
            if to_drop:
                dropped = min(to_drop, len(text))
                text, to_drop = text[dropped:], to_drop - dropped
            if not skip:
                out.append(text)
    flush()
    joined = "".join(out).encode("utf-16-le", "surrogatepass").decode("utf-16-le", "replace")
    return "\n".join(line.rstrip() for line in joined.split("\n")).strip("\n")


# ---------------------------------------------------------------------------- Mail


def _email_docs(msg: email.message.EmailMessage, rel: str, depth: int) -> list[Document]:
    """The message as one Document (headers + body), then one per attachment."""
    body_part = msg.get_body(preferencelist=("plain", "html"))
    attachments = _email_attachments(msg, body_part)
    header = [f"{h}: {msg[h]}" for h in ("Date", "From", "To", "Cc", "Subject") if msg[h]]
    if attachments:
        header.append("Attachments: " + ", ".join(name for name, _, _ in attachments))
    body = ""
    if body_part is not None:
        body = body_part.get_content()
        if body_part.get_content_type() == "text/html":
            body = _html_to_text(body)
    doc = Document(rel, "\n".join(header) + "\n\n" + body)
    if not body.strip() and not header:
        doc.status = "empty"
    docs = [doc]
    for name, content, ext in attachments:
        docs.extend(_attachment_docs(content, ext, f"{rel} > {name}", depth))
    return docs


def _email_attachments(msg, body_part) -> list[tuple[str, bytes | email.message.Message, str]]:
    """(unique name, bytes or attached email, extension) for every real attachment.
    Pictures shown inline in the body (logos, signatures) are left out silently."""
    found = []
    used: set[str] = set()
    for n, part in enumerate(_leaf_parts(msg), 1):
        if part is body_part:
            continue
        ctype = part.get_content_type()
        disposition = part.get_content_disposition()
        filename = part.get_filename()
        if ctype.startswith("image/") and disposition == "inline":
            continue
        if ctype == "message/rfc822":
            inner = part.get_payload(0) if part.is_multipart() else part.get_payload()
            subject = str(inner.get("Subject", "") or "") if hasattr(inner, "get") else ""
            name = _one_line(filename or f"{subject or 'attached message'}.eml")
            found.append((_unique(name, used), inner, ".eml"))
            continue
        if not filename and disposition != "attachment":
            continue  # alternative versions of the body (the HTML of a plain-text message)
        ext = Path(filename).suffix.lower() if filename else ""
        if not ext:
            ext = mimetypes.guess_extension(ctype) or ""
        name = filename or f"attachment {n}{ext}"
        found.append((_unique(_one_line(name), used), _part_bytes(part), ext))
    return found


def _leaf_parts(part) -> Iterator:
    """Non-multipart parts in order, without descending into attached emails."""
    if part.get_content_maintype() == "multipart":
        for sub in part.iter_parts():
            yield from _leaf_parts(sub)
    else:
        yield part


def _part_bytes(part) -> bytes:
    if part.get_content_maintype() == "text":
        # Text attachments: honour the declared charset, hand on UTF-8.
        with contextlib.suppress(Exception):
            return part.get_content().encode("utf-8")
    return part.get_payload(decode=True) or b""


def _attachment_docs(content, ext: str, rel: str, depth: int) -> list[Document]:
    if depth + 1 > MAX_NESTING:
        return [Document(rel, status="skipped: attachments nested too deeply")]
    if isinstance(content, email.message.Message):
        try:
            if not isinstance(content, email.message.EmailMessage):
                content = email.message_from_bytes(content.as_bytes(), policy=email.policy.default)
            return _email_docs(content, rel, depth + 1)
        except Exception as exc:
            return [Document(rel, status=f"error: {type(exc).__name__}: {exc}")]
    if isinstance(content, dict):  # an Outlook message embedded in a .msg
        try:
            return _msg_docs(content, rel, depth + 1)
        except Exception as exc:
            return [Document(rel, status=f"error: {type(exc).__name__}: {exc}")]
    return _read(content, ext, rel, depth + 1)


def _one_line(text: str, limit: int = 120) -> str:
    text = " ".join(str(text).split())
    return text if len(text) <= limit else text[: limit - 1] + "\u2026"


def _unique(name: str, used: set[str]) -> str:
    """Two attachments called 'scan.pdf' must not share a citation."""
    candidate, n = name, 2
    while candidate in used:
        stem, ext = os.path.splitext(name)
        candidate, n = f"{stem} ({n}){ext}", n + 1
    used.add(candidate)
    return candidate


def _read_mbox(src: Source, rel: str, depth: int) -> list[Document]:
    """Each message becomes '<mbox> > <n> <subject>' plus its attachments."""
    docs = []
    with _on_disk(src, ".mbox") as path:
        box = mailbox.mbox(
            path,
            factory=lambda f: email.message_from_binary_file(f, policy=email.policy.default),
            create=False,
        )
        try:
            for n, msg in enumerate(box, 1):
                subject = ""
                with contextlib.suppress(Exception):
                    subject = _one_line(msg.get("Subject", "") or "")
                name = f"{rel} > {n} {subject or '(no subject)'}"
                try:
                    docs.extend(_email_docs(msg, name, depth + 1))
                except Exception as exc:
                    docs.append(Document(name, status=f"error: {type(exc).__name__}: {exc}"))
        finally:
            box.close()
    return docs or [Document(rel, status="empty")]


# ------------------------------------------------------------------ Outlook .msg
#
# A .msg file is an OLE compound file. Each property of the message is a stream named
# "__substg1.0_IIIITTTT" (IIII = property id, TTTT = type) or, for fixed-size values such
# as dates, an entry in the "__properties_version1.0" stream. Attachments live in
# "__attach_version1.0_#N" storages; an attached Outlook message is a nested storage.
# _ole_message turns the file into a plain dict, _msg_docs turns that dict into Documents.

_MSG_SUBJECT, _MSG_BODY, _MSG_HTML, _MSG_RTF = "0037", "1000", "1013", "1009"
_MSG_SENDER_NAME, _MSG_HEADERS = "0C1A", "007D"
_MSG_SENDER_ADDRS = ("5D01", "0C1F", "0065")
_MSG_TO, _MSG_CC = "0E04", "0E03"
_MSG_DATES = ("0039", "0E06")  # client submit time, delivery time
_ATT_NAMES = ("3707", "3704", "3001")  # long filename, short filename, display name
_ATT_DATA, _ATT_MIME, _ATT_HIDDEN = "3701", "370E", "7FFE"


def _ole_message(data: bytes) -> dict:
    import olefile

    with olefile.OleFileIO(data) as ole:
        return _ole_object(ole, [], header=32)


def _ole_object(ole, prefix: list[str], header: int) -> dict:
    """Properties of the storage at prefix as {"props": {id: value}, "attachments": [...]}."""
    props: dict[str, object] = {}
    depth = len(prefix)
    for entry in ole.listdir(streams=True, storages=False):
        if len(entry) != depth + 1 or entry[:depth] != prefix:
            continue
        name = entry[-1]
        if name.startswith("__substg1.0_") and len(name) == 20:
            pid, ptype = name[12:16].upper(), name[16:20].upper()
            raw = ole.openstream(entry).read()
            if ptype == "001F":
                props[pid] = raw.decode("utf-16-le", errors="replace").rstrip("\x00")
            elif ptype == "001E":
                props[pid] = _decode(raw).rstrip("\x00")
            elif ptype == "0102":
                props[pid] = raw
    if ole.exists("/".join(prefix + ["__properties_version1.0"])):
        raw = ole.openstream(prefix + ["__properties_version1.0"]).read()
        props.update((k, v) for k, v in _fixed_props(raw[header:]).items() if k not in props)
    attachments = []
    for entry in ole.listdir(streams=False, storages=True):
        if len(entry) == depth + 1 and entry[:depth] == prefix and entry[-1].startswith("__attach_version1.0_"):
            attachments.append(_ole_attachment(ole, entry))
    return {"props": props, "attachments": attachments}


def _ole_attachment(ole, prefix: list[str]) -> dict:
    att = _ole_object(ole, prefix, header=8)
    embedded = prefix + ["__substg1.0_3701000D"]
    if ole.exists("/".join(embedded)) and (
        ole.exists("/".join(embedded + ["__properties_version1.0"]))
        or ole.exists("/".join(embedded + ["__substg1.0_001A001F"]))
        or ole.exists("/".join(embedded + ["__substg1.0_001A001E"]))
    ):
        att["message"] = _ole_object(ole, embedded, header=24)
    return att


def _fixed_props(raw: bytes) -> dict[str, object]:
    """Fixed-size properties: 16-byte entries of tag, flags and an 8-byte value."""
    props: dict[str, object] = {}
    for i in range(0, len(raw) - 15, 16):
        tag = int.from_bytes(raw[i : i + 4], "little")
        ptype, pid = tag & 0xFFFF, f"{tag >> 16:04X}"
        value = raw[i + 8 : i + 16]
        if ptype == 0x0040:  # time: 100 ns ticks since 1601
            ticks = int.from_bytes(value, "little")
            with contextlib.suppress(OverflowError, ValueError):
                epoch = dt.datetime(1601, 1, 1, tzinfo=dt.timezone.utc)
                props[pid] = epoch + dt.timedelta(microseconds=ticks // 10)
        elif ptype == 0x000B:  # boolean
            props[pid] = bool(value[0])
        elif ptype == 0x0003:  # 32-bit integer
            props[pid] = int.from_bytes(value[:4], "little", signed=True)
    return props


def _msg_docs(message: dict, rel: str, depth: int) -> list[Document]:
    """Documents for a parsed .msg: the message (headers and body) then its attachments."""
    props = message.get("props", {})
    names: set[str] = set()
    attachments = []
    for n, att in enumerate(message.get("attachments", []), 1):
        a = att.get("props", {})
        mime = str(a.get(_ATT_MIME, "") or "").lower()
        if mime.startswith("image/") and a.get(_ATT_HIDDEN) is True:
            continue  # inline picture in the body (logo, signature)
        filename = next((str(a[k]) for k in _ATT_NAMES if a.get(k)), "")
        if "message" in att:
            inner = att["message"].get("props", {})
            name = filename or f"{_one_line(str(inner.get(_MSG_SUBJECT, '')) or 'attached message')}.msg"
            attachments.append((_unique(_one_line(name), names), att["message"], ".msg"))
            continue
        ext = Path(filename).suffix.lower() if filename else (mimetypes.guess_extension(mime) or "")
        name = filename or f"attachment {n}{ext}"
        data = a.get(_ATT_DATA)
        attachments.append((_unique(_one_line(name), names), data if isinstance(data, bytes) else b"", ext))

    header = []
    date = next((props[k] for k in _MSG_DATES if isinstance(props.get(k), dt.datetime)), None)
    if date is not None:
        header.append(f"Date: {email.utils.format_datetime(date)}")
    else:
        m = re.search(r"^Date:[ \t]*(.+)$", str(props.get(_MSG_HEADERS, "")), re.MULTILINE)
        if m:
            header.append(f"Date: {m.group(1).strip()}")
    sender = str(props.get(_MSG_SENDER_NAME, "") or "")
    address = next((str(props[k]) for k in _MSG_SENDER_ADDRS if "@" in str(props.get(k, ""))), "")
    if address and address.lower() != sender.lower():
        sender = f"{sender} <{address}>" if sender else address
    for label, value in (("From", sender), ("To", props.get(_MSG_TO)), ("Cc", props.get(_MSG_CC)),
                         ("Subject", props.get(_MSG_SUBJECT))):
        if value:
            header.append(f"{label}: {value}")
    if attachments:
        header.append("Attachments: " + ", ".join(name for name, _, _ in attachments))

    body = props.get(_MSG_BODY)
    if not isinstance(body, str) or not body.strip():
        html = props.get(_MSG_HTML)
        if isinstance(html, bytes):
            html = _decode(html)
        if isinstance(html, str) and html.strip():
            body = _html_to_text(html)
        elif isinstance(props.get(_MSG_RTF), bytes):
            body = _rtf_to_text(_decompress_rtf(props[_MSG_RTF]))
        else:
            body = ""
    doc = Document(rel, "\n".join(header) + "\n\n" + body.replace("\r\n", "\n"))
    if not header and not body.strip():
        doc.status = "empty"
    docs = [doc]
    for name, content, ext in attachments:
        docs.extend(_attachment_docs(content, ext, f"{rel} > {name}", depth))
    return docs


_RTF_PREBUF = (
    b"{\\rtf1\\ansi\\mac\\deff0\\deftab720{\\fonttbl;}{\\f0\\fnil \\froman \\fswiss \\fmodern "
    b"\\fscript \\fdecor MS Sans SerifSymbolArialTimes New RomanCourier{\\colortbl\\red0\\green0"
    b"\\blue0\r\n\\par \\pard\\plain\\f0\\fs20\\b\\i\\u\\tab\\tx"
)


def _decompress_rtf(data: bytes) -> bytes:
    """Outlook's compressed RTF body (MS-OXRTFCP): LZ77 against a 4 KB dictionary."""
    if len(data) < 16:
        return b""
    comp_size, raw_size = int.from_bytes(data[0:4], "little"), int.from_bytes(data[4:8], "little")
    magic = data[8:12]
    if magic == b"MELA":  # stored uncompressed
        return data[16 : 16 + raw_size]
    if magic != b"LZFu":
        raise ValueError("unknown compressed RTF format")
    window = bytearray(_RTF_PREBUF.ljust(4096, b"\x00"))
    write = len(_RTF_PREBUF)
    out = bytearray()
    pos, end = 16, min(len(data), comp_size + 4)  # comp_size counts from byte 4
    while pos < end:
        flags = data[pos]
        pos += 1
        for bit in range(8):
            if pos >= end:
                break
            if flags & (1 << bit):
                if pos + 1 >= len(data):
                    return bytes(out)
                ref = int.from_bytes(data[pos : pos + 2], "big")
                pos += 2
                offset, length = ref >> 4, (ref & 0xF) + 2
                if offset == write:
                    return bytes(out)
                for k in range(length):
                    byte = window[(offset + k) % 4096]
                    out.append(byte)
                    window[write] = byte
                    write = (write + 1) % 4096
            else:
                byte = data[pos]
                pos += 1
                out.append(byte)
                window[write] = byte
                write = (write + 1) % 4096
    return bytes(out)


# --------------------------------------------------------------------------- HTML

_CELL = "\x1f"  # marks the end of a table cell until lines are tidied
# Inline styles that hide text from a reader (CSS classes can't be resolved here).
_HIDING_STYLE = re.compile(
    r"display\s*:\s*none|visibility\s*:\s*(hidden|collapse)"
    r"|(font-size|opacity)\s*:\s*(0+(\.0*)?|\.0+)\s*([a-z%]+)?\s*(;|!|$)",
    re.IGNORECASE,
)
_VOID = {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "source", "track", "wbr"}
# Tags whose end tag may be left out: a new sibling of the same kind closes the old one.
_OPTIONAL_END = {"p", "li", "td", "th", "tr", "dt", "dd", "option"}


class _TextExtractor(HTMLParser):
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "table", "section", "article"}
    NEVER_SHOWN = {"script", "style", "template", "noscript"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._hide_tag: str | None = None  # element being skipped, and how deep we are in it
        self._hide_depth = 0

    def handle_starttag(self, tag, attrs):
        if self._hide_tag is not None:
            if tag != self._hide_tag:
                return
            if tag in _OPTIONAL_END and self._hide_depth == 1:
                self._hide_tag, self._hide_depth = None, 0  # sibling closes the hidden one
            else:
                self._hide_depth += 1
                return
        if tag not in _VOID and (tag in self.NEVER_SHOWN or _hidden(dict(attrs))):
            # Text a reader can't see must not reach the model: it could carry instructions.
            self._hide_tag, self._hide_depth = tag, 1
            return
        if tag in self.BLOCK:
            self.parts.append("\n")

    def handle_startendtag(self, tag, attrs):
        if self._hide_tag is None and tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if self._hide_tag is not None:
            if tag == self._hide_tag:
                self._hide_depth -= 1
                if self._hide_depth == 0:
                    self._hide_tag = None
            return
        if tag in ("td", "th"):
            self.parts.append(_CELL)
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if self._hide_tag is None:
            self.parts.append(data)


def _hidden(attrs: dict) -> bool:
    if "hidden" in attrs or (attrs.get("aria-hidden") or "").strip().lower() == "true":
        return True
    return bool(_HIDING_STYLE.search(attrs.get("style") or ""))


def _html_to_text(html: str) -> str:
    parser = _TextExtractor()
    parser.feed(html)
    parser.close()
    lines = []
    for line in "".join(parser.parts).splitlines():
        cells = [" ".join(cell.split()) for cell in line.split(_CELL)]
        tidy = "\t".join(cells).strip()
        if tidy:
            lines.append(tidy)
    return "\n".join(lines)
