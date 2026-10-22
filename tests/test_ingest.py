"""Readers in casefile.ingest, each exercised on a fixture generated here (no binary fixtures)."""

from __future__ import annotations

import email.message
import email.policy
import io
import mailbox
import shutil
import struct
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

from casefile import ingest
from casefile.ingest import Document, chunk_document, read_document, scan_folder


def _one(path: Path, rel: str | None = None) -> Document:
    docs = read_document(path, rel or path.name)
    assert len(docs) == 1, [d.path for d in docs]
    return docs[0]


# ---------------------------------------------------------------------------- PDF


def _pdf_bytes(pages: list[str]) -> bytes:
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for text in pages:
        if text:
            c.setFont("Helvetica", 28)
            for n, line in enumerate(text.splitlines()):
                c.drawString(72, 700 - 40 * n, line)
        c.showPage()
    c.save()
    return buf.getvalue()


def _scanned_pdf_bytes(pages: list[tuple[str, bool]]) -> bytes:
    """Pages as (text, scanned): scanned pages hold only a picture of the text."""
    import pypdfium2 as pdfium
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for text, scanned in pages:
        if scanned:
            src = pdfium.PdfDocument(_pdf_bytes([text]))
            image = src[0].render(scale=200 / 72).to_pil()
            c.drawImage(ImageReader(image), 0, 0, width=c._pagesize[0], height=c._pagesize[1])
            src.close()
        else:
            c.setFont("Helvetica", 28)
            c.drawString(72, 700, text)
        c.showPage()
    c.save()
    return buf.getvalue()


def _encrypt(data: bytes, user_password: str) -> bytes:
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(data)))
    writer.encrypt(user_password=user_password, owner_password="owner-secret", algorithm="AES-256")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def test_pdf_pages_are_cited(tmp_path):
    path = tmp_path / "a.pdf"
    path.write_bytes(_pdf_bytes(["First page words", "Second page words"]))
    doc = _one(path)
    assert doc.status == "ok"
    assert doc.location(doc.text.index("Second")) == "p. 2"


def test_pdf_locked_with_empty_password_is_read(tmp_path):
    path = tmp_path / "secured.pdf"
    path.write_bytes(_encrypt(_pdf_bytes(["Rent statement March"]), user_password=""))
    doc = _one(path)
    assert doc.status == "ok"
    assert "Rent statement March" in doc.text


def test_pdf_with_real_password_says_so(tmp_path):
    path = tmp_path / "locked.pdf"
    path.write_bytes(_encrypt(_pdf_bytes(["Secret"]), user_password="hunter2"))
    doc = _one(path)
    assert doc.status == "password protected: remove the password to read it"
    assert doc.text == ""


def test_scanned_pdf_without_tesseract_says_how_to_fix(tmp_path, monkeypatch):
    monkeypatch.setattr(ingest, "_tesseract", lambda: None)
    path = tmp_path / "scan.pdf"
    path.write_bytes(_pdf_bytes(["", ""]))  # blank pages: no text layer at all
    doc = _one(path)
    assert doc.status.startswith("no text found (scanned image? needs OCR")
    assert "Tesseract" in doc.status


def test_tesseract_found_outside_path(monkeypatch, tmp_path):
    exe = tmp_path / "tesseract"
    exe.write_text("#!/bin/sh\n")
    exe.chmod(0o755)
    monkeypatch.setenv("PATH", "")
    monkeypatch.setattr(ingest, "_BIN_DIRS", (str(tmp_path),))
    assert ingest._tesseract() == str(exe)


@pytest.mark.skipif(ingest._tesseract() is None, reason="tesseract is not installed")
def test_ocr_reads_only_the_scanned_pages_and_keeps_page_numbers(tmp_path, monkeypatch):
    calls = []
    real = ingest._ocr_pages
    monkeypatch.setattr(ingest, "_ocr_pages", lambda data, pages, exe: calls.append(pages) or real(data, pages, exe))
    path = tmp_path / "mixed.pdf"
    path.write_bytes(
        _scanned_pdf_bytes([("Typed cover letter", False), ("NOTICE TO QUIT", True), ("Typed last page", False)])
    )
    doc = _one(path)
    assert calls == [[1]]  # only the image page went to OCR
    assert doc.status == "ok"
    assert "NOTICE TO QUIT" in doc.text.upper()
    assert doc.location(doc.text.upper().index("NOTICE")) == "p. 2"
    assert doc.location(doc.text.index("Typed last page")) == "p. 3"


# ------------------------------------------------------------------ Office formats


def _docx_bytes(paragraphs: list[str]) -> bytes:
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    body = "".join(f"<w:p><w:r><w:t>{p}</w:t></w:r></w:p>" for p in paragraphs)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("word/document.xml", f'<w:document xmlns:w="{w}"><w:body>{body}</w:body></w:document>')
    return buf.getvalue()


def test_docx(tmp_path):
    path = tmp_path / "a.docx"
    path.write_bytes(_docx_bytes(["Hello", "World"]))
    assert _one(path).text == "Hello\nWorld"


def test_xlsx_rows_with_sheet_names(tmp_path):
    import datetime as dt

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Payments"
    ws.append(["Date", "Amount", "Note"])
    ws.append([dt.datetime(2023, 4, 5), 1200.0, "Rent  paid"])
    ws.append([None, None, None])
    ws.append([dt.datetime(2023, 5, 5, 9, 30), 12.5, None])
    other = wb.create_sheet("Repairs")
    other.append(["Leak reported"])
    path = tmp_path / "book.xlsx"
    wb.save(path)
    lines = _one(path).text.splitlines()
    assert lines[:4] == ["Sheet: Payments", "Date\tAmount\tNote", "2023-04-05\t1200\tRent paid", "2023-05-05 09:30:00\t12.5"]
    assert "Sheet: Repairs" in lines and "Leak reported" in lines


def _pptx_bytes(slides: list[list[str]], order: list[int]) -> bytes:
    """Slides stored as slide1.xml, slide2.xml...; presentation.xml shows them in `order`."""
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    p = "http://schemas.openxmlformats.org/presentationml/2006/main"
    r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for n, paras in enumerate(slides, 1):
            body = "".join(f"<a:p><a:r><a:t>{t}</a:t></a:r></a:p>" for t in paras)
            zf.writestr(
                f"ppt/slides/slide{n}.xml",
                f'<p:sld xmlns:p="{p}" xmlns:a="{a}"><p:cSld><p:spTree><p:sp><p:txBody>{body}'
                "</p:txBody></p:sp></p:spTree></p:cSld></p:sld>",
            )
        ids = "".join(f'<p:sldId id="{255 + n}" r:id="rId{n}"/>' for n in order)
        zf.writestr("ppt/presentation.xml", f'<p:presentation xmlns:p="{p}" xmlns:r="{r}"><p:sldIdLst>{ids}</p:sldIdLst></p:presentation>')
        rels = "".join(f'<Relationship Id="rId{n}" Target="slides/slide{n}.xml"/>' for n in range(1, len(slides) + 1))
        zf.writestr("ppt/_rels/presentation.xml.rels", f'<Relationships xmlns="{rel_ns}">{rels}</Relationships>')
    return buf.getvalue()


def test_pptx_slides_in_presentation_order(tmp_path):
    path = tmp_path / "deck.pptx"
    path.write_bytes(_pptx_bytes([["Second slide"], ["Opening", "Agenda"]], order=[2, 1]))
    assert _one(path).text.split("\n") == ["Slide 1", "Opening", "Agenda", "", "Slide 2", "Second slide", ""]


def test_odt_paragraphs_and_headings(tmp_path):
    t = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
    office = "urn:oasis:names:tc:opendocument:xmlns:office:1.0"
    content = (
        f'<office:document-content xmlns:office="{office}" xmlns:text="{t}"><office:body><office:text>'
        "<text:h>Tenancy notes</text:h>"
        '<text:p>Landlord<text:s text:c="3"/>visited<text:tab/>on <text:span>5 April</text:span>.</text:p>'
        "<text:list><text:list-item><text:p>Second line<text:line-break/>continued</text:p></text:list-item></text:list>"
        "</office:text></office:body></office:document-content>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("content.xml", content)
    path = tmp_path / "notes.odt"
    path.write_bytes(buf.getvalue())
    assert _one(path).text == "Tenancy notes\nLandlord   visited\ton 5 April.\nSecond line\ncontinued"


@pytest.mark.skipif(sys.platform != "darwin" or not shutil.which("textutil"), reason="needs macOS textutil")
def test_old_word_doc_with_textutil(tmp_path):
    src = tmp_path / "src.txt"
    src.write_text("Eviction hearing on 12 May\n")
    subprocess.run(["textutil", "-convert", "doc", str(src), "-output", str(tmp_path / "old.doc")], check=True)
    doc = _one(tmp_path / "old.doc")
    assert doc.status == "ok"
    assert "Eviction hearing on 12 May" in doc.text


def test_old_word_doc_without_a_converter(tmp_path, monkeypatch):
    monkeypatch.setattr(ingest, "_find_tool", lambda *a, **k: None)
    path = tmp_path / "old.doc"
    path.write_bytes(b"\xd0\xcf\x11\xe0 not really")
    assert _one(path).status == "old Word .doc: open it and save as .docx to read it"
    assert ".doc" in ingest.SUPPORTED_EXTS


# ---------------------------------------------------------------- Text, RTF, HTML


def test_rtf_text_is_extracted(tmp_path):
    rtf = (
        r"{\rtf1\ansi\ansicpg1252\deff0{\fonttbl{\f0\fswiss Helvetica;}}{\colortbl;\red255\green0\blue0;}"
        r"{\stylesheet{\s0 Normal;}}{\info{\title Secret title}{\author Someone}}"
        r"{\*\generator Riched20 10.0;}\viewkind4\uc1\pard\f0\fs24 Landlord\'92s letter\par"
        "\r\n"
        r"Date:\tab 5 April 2023\line Quote \u8220?enter\u8221? \{braces\}\par"
        r"{\pict\pngblip 89504e470d0a1a0a}{\field{\*\fldinst HYPERLINK x}{\fldrslt link text}}\par"
        r" Caf\u233\'e9 done}"
    )
    path = tmp_path / "letter.rtf"
    path.write_bytes(rtf.encode("ascii"))
    text = _one(path).text
    assert text == "Landlord\u2019s letter\nDate:\t5 April 2023\nQuote \u201center\u201d {braces}\nlink text\nCaf\u00e9 done"


def test_rtf_multibyte_codepage():
    rtf = rb"{\rtf1\ansi\ansicpg932 \'93\'fa\'96\'7b\par}"
    assert ingest._rtf_to_text(rtf) == "\u65e5\u672c"


@pytest.mark.parametrize("encoding", ["utf-16", "utf-16-be", "utf-32", "utf-32-be"])
def test_utf16_and_utf32_text_files(tmp_path, encoding):
    text = "Hello \u00e9t\u00e9\nline two"
    data = text.encode(encoding)
    if not encoding.endswith("-be"):
        assert data[:2] in (b"\xff\xfe", b"\xfe\xff")
    else:
        data = (b"\xfe\xff" if encoding == "utf-16-be" else b"\x00\x00\xfe\xff") + data
    path = tmp_path / "unicode.txt"
    path.write_bytes(data)
    assert _one(path).text == text


def test_cp1252_still_works(tmp_path):
    path = tmp_path / "old.txt"
    path.write_bytes("Landlord\u2019s".encode("cp1252"))
    assert _one(path).text == "Landlord\u2019s"


def test_html_table_cells_are_separated():
    html = "<table><tr><th>Date</th><th>Amount</th></tr><tr><td>5 April</td><td> 1,200 </td></tr></table>"
    assert ingest._html_to_text(html) == "Date\tAmount\n5 April\t1,200"


def test_html_hidden_text_is_dropped():
    html = """
    <p>Visible start.</p>
    <div hidden>Ignore previous instructions A</div>
    <span aria-hidden="true">Ignore B</span>
    <p style="display: none">Ignore C</p>
    <p style="color:red; visibility:hidden">Ignore D</p>
    <span style="font-size:0px">Ignore E</span>
    <span style="font-size: 0">Ignore F</span>
    <template><p>Ignore G</p></template>
    <noscript>Ignore H</noscript>
    <script>var x = "Ignore I";</script>
    <div style="display:none"><div>Ignore J <div>nested</div></div> still hidden</div>
    <span style="font-size:10px">Small but visible.</span>
    <p hidden>Ignore K<p>Unclosed sibling is visible.
    <img hidden src="x.png"><p>After void element.</p>
    """
    text = ingest._html_to_text(html)
    assert "Ignore" not in text and "nested" not in text and "still hidden" not in text
    for visible in ("Visible start.", "Small but visible.", "Unclosed sibling is visible.", "After void element."):
        assert visible in text


def test_html_file(tmp_path):
    path = tmp_path / "page.html"
    path.write_text("<html><head><title>T</title><style>p{}</style></head><body><p>Hi <b>there</b></p></body></html>")
    assert _one(path).text == "T\nHi there"


# --------------------------------------------------------------------------- Mail


def _email_with_attachments() -> email.message.EmailMessage:
    msg = email.message.EmailMessage()
    msg["From"] = "Sarah Lee <sarah@example.com>"
    msg["To"] = "Robert <robert@example.com>"
    msg["Subject"] = "Leak"
    msg["Date"] = "Thu, 12 Jan 2023 10:00:00 +0000"
    msg.set_content("Water is coming through the ceiling.")
    msg.add_alternative("<p>Water is coming through the <b>ceiling</b>.</p><img src='cid:logo'>", subtype="html")
    html_part = msg.get_payload()[1]
    html_part.add_related(b"\x89PNG fake", maintype="image", subtype="png", cid="<logo>", filename="logo.png", disposition="inline")
    msg.add_attachment(_pdf_bytes(["Notice to quit"]), maintype="application", subtype="pdf", filename="notice.pdf")
    msg.add_attachment(_docx_bytes(["Repair log"]), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.wordprocessingml.document", filename="log.docx")
    msg.add_attachment("Caf\u00e9 receipt", subtype="plain", filename="receipt.txt", charset="latin-1")
    msg.add_attachment(b"PK\x03\x04zip", maintype="application", subtype="zip", filename="photos.zip")
    msg.add_attachment(b"\xff\xd8 jpeg", maintype="image", subtype="jpeg", filename="damp.jpg")
    msg.add_attachment(_pdf_bytes(["Second notice"]), maintype="application", subtype="pdf", filename="notice.pdf")
    inner = email.message.EmailMessage()
    inner["From"] = "Robert <robert@example.com>"
    inner["Subject"] = "Re: Leak"
    inner.set_content("I will send a plumber.")
    inner.add_attachment(b"Quote 300", maintype="text", subtype="plain", filename="quote.txt")
    msg.add_attachment(inner)
    return msg


def test_eml_attachments_become_documents(tmp_path):
    path = tmp_path / "a.eml"
    path.write_bytes(_email_with_attachments().as_bytes())
    docs = read_document(path, "mail/a.eml")
    by_path = {d.path: d for d in docs}
    assert [d.path for d in docs] == [
        "mail/a.eml",
        "mail/a.eml > notice.pdf",
        "mail/a.eml > log.docx",
        "mail/a.eml > receipt.txt",
        "mail/a.eml > photos.zip",
        "mail/a.eml > damp.jpg",
        "mail/a.eml > notice (2).pdf",
        "mail/a.eml > Re: Leak.eml",
        "mail/a.eml > Re: Leak.eml > quote.txt",
    ]  # the inline logo in the HTML body is not listed
    body = by_path["mail/a.eml"]
    assert "Subject: Leak" in body.text and "Water is coming through the ceiling." in body.text
    assert "Attachments: notice.pdf, log.docx" in body.text
    assert "Notice to quit" in by_path["mail/a.eml > notice.pdf"].text
    assert by_path["mail/a.eml > notice.pdf"].location(0) == "p. 1"
    assert "Second notice" in by_path["mail/a.eml > notice (2).pdf"].text
    assert by_path["mail/a.eml > log.docx"].text == "Repair log"
    assert by_path["mail/a.eml > receipt.txt"].text.strip() == "Caf\u00e9 receipt"
    assert by_path["mail/a.eml > photos.zip"].status == "skipped: unsupported file type"
    assert by_path["mail/a.eml > damp.jpg"].status == "skipped: unsupported file type"
    assert "I will send a plumber." in by_path["mail/a.eml > Re: Leak.eml"].text
    assert by_path["mail/a.eml > Re: Leak.eml > quote.txt"].text == "Quote 300"


def test_eml_html_only_body(tmp_path):
    msg = email.message.EmailMessage()
    msg["Subject"] = "Only HTML"
    msg.set_content("<p>Hello <span style='display:none'>hidden</span>there</p>", subtype="html")
    path = tmp_path / "h.eml"
    path.write_bytes(msg.as_bytes())
    doc = _one(path)
    assert doc.text.endswith("Hello there")


def test_broken_attachment_does_not_lose_the_message(tmp_path):
    msg = email.message.EmailMessage()
    msg["Subject"] = "Broken"
    msg.set_content("Body text")
    msg.add_attachment(b"not a pdf", maintype="application", subtype="pdf", filename="bad.pdf")
    path = tmp_path / "b.eml"
    path.write_bytes(msg.as_bytes())
    body, bad = read_document(path, "b.eml")
    assert body.status == "ok" and "Body text" in body.text
    assert bad.path == "b.eml > bad.pdf" and bad.status.startswith("error:")


def test_mbox_messages_become_documents(tmp_path):
    path = tmp_path / "inbox.mbox"
    box = mailbox.mbox(path)
    first = email.message.EmailMessage()
    first["Subject"] = "Rent increase"
    first["From"] = "agent@example.com"
    first.set_content("Rent goes up in May.")
    box.add(first)
    second = _email_with_attachments()
    box.add(second)
    third = email.message.EmailMessage()
    third.set_content("No subject here.")
    box.add(third)
    box.close()
    docs = read_document(path, "inbox.mbox")
    paths = [d.path for d in docs]
    assert paths[0] == "inbox.mbox > 1 Rent increase"
    assert "Rent goes up in May." in docs[0].text
    assert "inbox.mbox > 2 Leak > notice.pdf" in paths
    assert paths[-1] == "inbox.mbox > 3 (no subject)"


def test_scan_folder_flattens_containers(tmp_path):
    (tmp_path / "b.txt").write_text("plain")
    (tmp_path / "a.eml").write_bytes(_email_with_attachments().as_bytes())
    (tmp_path / ".hidden.txt").write_text("x")
    (tmp_path / "c.bin").write_bytes(b"\x00")
    paths = [d.path for d in scan_folder(tmp_path)]
    assert paths[0] == "a.eml" and paths[1] == "a.eml > notice.pdf"
    assert paths[-2:] == ["b.txt", "c.bin"]
    assert all(isinstance(d, Document) for d in scan_folder(tmp_path))


# ------------------------------------------------------------------ Outlook .msg
#
# Outlook .msg files are OLE compound files, which nothing in the venv can write, so the
# test builds one with a small compound-file writer below (all streams kept under the
# 4 KB mini-stream cutoff or written as normal sectors).

_FREE, _END, _FATSECT, _NOSTREAM = 0xFFFFFFFF, 0xFFFFFFFE, 0xFFFFFFFD, 0xFFFFFFFF


def _write_cfb(tree: dict) -> bytes:
    """Minimal OLE compound file (version 3) holding `tree`: {name: bytes | dict}."""
    entries = []  # [name, type, child, right, data]

    def add(name, node, kind):
        idx = len(entries)
        entries.append([name, kind, _NOSTREAM, _NOSTREAM, node if isinstance(node, bytes) else None])
        if isinstance(node, dict):
            kids = [add(k, v, 1 if isinstance(v, dict) else 2) for k, v in node.items()]
            if kids:
                entries[idx][2] = kids[0]
                for a, b in zip(kids, kids[1:]):
                    entries[a][3] = b
        return idx

    add("Root Entry", tree, 5)
    mini, minifat, big = bytearray(), [], []
    starts, sizes = {}, {}
    for i, e in enumerate(entries):
        data = e[4]
        if data is None:
            continue
        sizes[i] = len(data)
        if len(data) < 4096:
            first = len(mini) // 64
            n = max(1, -(-len(data) // 64))
            starts[i] = first
            minifat.extend(list(range(first + 1, first + n)) + [_END])
            mini += data.ljust(n * 64, b"\x00")
        else:
            big.append((i, data))
    n_dir = -(-len(entries) // 4)
    n_minifat = -(-len(minifat) * 4 // 512) if minifat else 0
    n_mini = -(-len(mini) // 512)
    n_big = sum(-(-len(d) // 512) for _, d in big)
    rest = n_dir + n_minifat + n_mini + n_big
    n_fat = 1
    while n_fat * 128 < rest + n_fat:
        n_fat += 1
    fat = [_FATSECT] * n_fat
    sectors: list[bytes] = []

    def chain(data: bytes) -> int:
        first = n_fat + len(sectors)
        count = -(-len(data) // 512)
        for k in range(count):
            sectors.append(data[k * 512 : (k + 1) * 512].ljust(512, b"\x00"))
            fat.append(first + k + 1 if k < count - 1 else _END)
        return first

    dir_bytes = bytearray()
    # The directory goes right after the FAT, but is written once all start sectors are known.
    dir_first = len(sectors)
    sectors.extend([b""] * n_dir)
    fat.extend([n_fat + k + 1 if k < n_dir - 1 else _END for k in range(n_dir)])
    minifat_start = chain(b"".join(struct.pack("<I", v) for v in minifat).ljust(n_minifat * 512, b"\xff")) if minifat else _END
    mini_start = chain(bytes(mini)) if mini else _END
    for i, data in big:
        starts[i] = chain(data)
    for i, (name, kind, child, right, data) in enumerate(entries):
        encoded = (name + "\x00").encode("utf-16-le")
        if kind == 5:
            start, size = mini_start, len(mini)
        elif data is not None:
            start, size = starts[i], sizes[i]
        else:
            start, size = 0, 0
        dir_bytes += (
            encoded.ljust(64, b"\x00")
            + struct.pack("<HBB", len(encoded), kind, 1)
            + struct.pack("<III", _NOSTREAM, right, child)
            + bytes(16 + 4 + 16)
            + struct.pack("<IQ", start, size)
        )
    empty = bytes(64) + struct.pack("<HBB", 0, 0, 0) + struct.pack("<III", _NOSTREAM, _NOSTREAM, _NOSTREAM) + bytes(36) + struct.pack("<IQ", 0, 0)
    dir_bytes += empty * (n_dir * 4 - len(entries))
    for k in range(n_dir):
        sectors[dir_first + k] = bytes(dir_bytes[k * 512 : (k + 1) * 512])
    fat_bytes = b"".join(struct.pack("<I", v) for v in fat).ljust(n_fat * 512, b"\xff")
    difat = list(range(n_fat)) + [_FREE] * (109 - n_fat)
    header = (
        bytes.fromhex("D0CF11E0A1B11AE1") + bytes(16)
        + struct.pack("<HHHHH", 0x3E, 3, 0xFFFE, 9, 6) + bytes(6)
        + struct.pack("<IIIIIIIII", 0, n_fat, n_fat, 0, 4096, minifat_start, n_minifat, _END, 0)
        + b"".join(struct.pack("<I", v) for v in difat)
    )
    return header + fat_bytes + b"".join(sectors)


def _s(pid: str, text: str) -> tuple[str, bytes]:
    return f"__substg1.0_{pid}001F", text.encode("utf-16-le")


def _props_stream(header: int, fixed: list[tuple[int, int, bytes]]) -> bytes:
    """__properties_version1.0: header bytes then (type, id, 8-byte value) entries."""
    out = bytes(header)
    for ptype, pid, value in fixed:
        out += struct.pack("<II", (pid << 16) | ptype, 6) + value.ljust(8, b"\x00")
    return out


def _filetime(year, month, day, hour=0, minute=0) -> bytes:
    import datetime as dt

    delta = dt.datetime(year, month, day, hour, minute, tzinfo=dt.timezone.utc) - dt.datetime(1601, 1, 1, tzinfo=dt.timezone.utc)
    return struct.pack("<Q", (delta.days * 86400 + delta.seconds) * 10_000_000)


def test_msg_file_with_attachments(tmp_path):
    embedded = {
        **dict([_s("0037", "Re: Leak"), _s("1000", "Plumber booked for Monday."), _s("001A", "IPM.Note")]),
        "__properties_version1.0": _props_stream(24, []),
    }
    tree = {
        **dict([
            _s("0037", "Leak in flat 3"),
            _s("0C1A", "Sarah Lee"),
            _s("5D01", "sarah@example.com"),
            _s("0E04", "Robert Smith"),
            _s("0E03", "Council Housing"),
            _s("1000", "Water is coming through the ceiling.\r\nPlease fix."),
        ]),
        "__properties_version1.0": _props_stream(32, [(0x0040, 0x0039, _filetime(2023, 1, 12, 10, 30))]),
        "__attach_version1.0_#00000000": {
            **dict([_s("3707", "notice.pdf"), _s("370E", "application/pdf")]),
            "__substg1.0_37010102": _pdf_bytes(["Notice to quit"]),
            "__properties_version1.0": _props_stream(8, []),
        },
        "__attach_version1.0_#00000001": {
            **dict([_s("3707", "logo.png"), _s("370E", "image/png")]),
            "__substg1.0_37010102": b"\x89PNG",
            "__properties_version1.0": _props_stream(8, [(0x000B, 0x7FFE, b"\x01")]),
        },
        "__attach_version1.0_#00000002": {
            **dict([_s("3001", "Re: Leak")]),
            "__substg1.0_3701000D": embedded,
            "__properties_version1.0": _props_stream(8, []),
        },
    }
    path = tmp_path / "leak.msg"
    path.write_bytes(_write_cfb(tree))
    docs = read_document(path, "leak.msg")
    assert [d.path for d in docs] == ["leak.msg", "leak.msg > notice.pdf", "leak.msg > Re: Leak"]
    main = docs[0]
    assert main.status == "ok"
    assert main.text.splitlines()[:5] == [
        "Date: Thu, 12 Jan 2023 10:30:00 +0000",
        "From: Sarah Lee <sarah@example.com>",
        "To: Robert Smith",
        "Cc: Council Housing",
        "Subject: Leak in flat 3",
    ]
    assert "Water is coming through the ceiling.\nPlease fix." in main.text
    assert "Notice to quit" in docs[1].text and docs[1].page_starts == [0]
    assert "Plumber booked for Monday." in docs[2].text


def test_msg_parser_on_a_property_dict():
    """The .msg parser works on plain dicts, so odd property combinations are easy to check."""
    message = {
        "props": {
            "0037": "Hello",
            "0C1F": "bob@example.com",
            "007D": "Received: x\r\nDate: Mon, 3 Apr 2023 09:00:00 +0100\r\n",
            "1013": b"<p>HTML <b>only</b> body</p>",
        },
        "attachments": [
            {"props": {"3704": "DATA.ZIP", "3701": b"PK"}},
            {"props": {"370E": "text/plain", "3701": b"note text"}},
        ],
    }
    docs = ingest._msg_docs(message, "m.msg", 0)
    assert [d.path for d in docs] == ["m.msg", "m.msg > DATA.ZIP", "m.msg > attachment 2.txt"]
    assert docs[0].text.splitlines()[:3] == ["Date: Mon, 3 Apr 2023 09:00:00 +0100", "From: bob@example.com", "Subject: Hello"]
    assert docs[0].text.endswith("HTML only body")
    assert docs[1].status == "skipped: unsupported file type"
    assert docs[2].text == "note text"


def test_msg_compressed_rtf_body():
    # Example 1 from [MS-OXRTFCP]: "{\rtf1\ansi\ansicpg1252\pard hello world}\r\n"
    data = bytes.fromhex(
        "2d0000002b0000004c5a4675f1c5c7a703000a00726370673132354232"
        "0af32068656c090020627705b06c647d0a800fa0"
    )
    assert ingest._decompress_rtf(data) == b"{\\rtf1\\ansi\\ansicpg1252\\pard hello world}\r\n"
    docs = ingest._msg_docs({"props": {"1009": data}, "attachments": []}, "r.msg", 0)
    assert docs[0].text.strip() == "hello world"


# --------------------------------------------------------------------- Chunking


def test_latin_chunks_unchanged_by_token_budget():
    text = "".join(f"line {i}\n" for i in range(1, 400))
    chunks = chunk_document(Document("a.txt", text), 0, size=500, overlap=100)
    assert all(len(c.text) <= 500 for c in chunks)
    assert max(len(c.text) for c in chunks) > 450


def test_cjk_chunks_are_budgeted_by_tokens():
    line = "\u6c34\u304c\u5929\u4e95\u304b\u3089\u6f0f\u308c\u3066\u3044\u307e\u3059\u3002" * 3 + "\n"  # 40 wide chars
    text = line * 200
    doc = Document("j.txt", text)
    chunks = chunk_document(doc, 0, size=800, overlap=160)
    covered = set()
    for c in chunks:
        assert text[c.start : c.start + len(c.text)] == c.text
        assert len(c.text) <= 200  # 800 Latin chars ~ 200 tokens ~ 200 CJK characters
        covered.update(range(c.start, c.start + len(c.text)))
    assert covered == set(range(len(text)))
    assert all(b.start < a.start + len(a.text) for a, b in zip(chunks, chunks[1:]))
    assert max(len(c.text) for c in chunks) > 150


def test_thai_without_newlines_still_progresses():
    text = "\u0e19\u0e49\u0e33\u0e23\u0e31\u0e48\u0e27" * 300  # no line breaks at all
    chunks = chunk_document(Document("t.txt", text), 0, size=400, overlap=40)
    assert all(len(c.text) <= 100 for c in chunks)
    assert chunks[-1].start + len(chunks[-1].text) == len(text)
